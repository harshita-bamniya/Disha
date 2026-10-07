"""Import the research sheets of UPSC_Master_Database_v2.xlsx into the reference tables.

    python -m app.scripts.import_reference_sheets <workbook.xlsx> --dry-run
    python -m app.scripts.import_reference_sheets <workbook.xlsx>
    python -m app.scripts.import_reference_sheets <workbook.xlsx> --verify     # compare the DB to the workbook

Sheets: Skill Framework, Salary Benchmarks, Industry Analysis, Candidate Segmentation, UPSC Ecosystem & Skills.
(The 'Master Database' sheet has its own importer: import_market_roles.)

All-or-nothing: every sheet is parsed and validated first; if ANY row of ANY sheet has a problem nothing is
written. Re-running is safe: rows are matched on their natural key and updated in place.
"""
import sys
from pathlib import Path

from app.modules.reference_data import parser as P

SOURCE = "UPSC_Master_Database_v2"

# natural key per model (attribute names)
KEYS = {
    "UpscSkillFramework": ("skill_category", "upsc_activity"),
    "SalaryBenchmark": ("role_category", "specific_role"),
    "IndustryAnalysis": ("industry",),
    "CandidateSegment": ("category_no",),
    "UpscEcosystemEntry": ("section", "position"),
}
# source header -> model attribute, for the independent verification (raw text columns only)
VERIFY = {
    "Skill Framework": ("UpscSkillFramework", ["skill_category", "upsc_activity", "skill_developed", "corporate_translation",
                                               "relevant_job_roles", "skill_depth", "industry_applicability"]),
    "Salary Benchmarks": ("SalaryBenchmark", ["role_category", "specific_role", "entry_text", "mid_text", "senior_text", "notes"]),
    "Industry Analysis": ("IndustryAnalysis", ["rank_no", "industry", "why_upsc_fit", "top_roles", "hiring_potential",
                                               "growth_trajectory", "primary_upsc_skills", "overall_fit_score"]),
    "Candidate Segmentation": ("CandidateSegment", ["category_no", "category_name", "knowledge_depth", "analytical_ability",
                                                    "research_capability", "communication", "strategic_thinking",
                                                    "problem_solving", "leadership_traits", "professional_maturity",
                                                    "employability_strengths", "employability_limitations", "overall_score"]),
}


def read_sheets(path: str) -> dict[str, list[tuple[int, tuple]]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {}
    for name, _parser, _model in P.SHEETS:
        if name not in wb.sheetnames:
            raise SystemExit(f"Sheet '{name}' not found. Sheets: {wb.sheetnames}")
        out[name] = [(i, tuple(r)) for i, r in enumerate(wb[name].iter_rows(values_only=True), 1)]
    return out


def parse_all(sheets):
    parsed, errors = {}, []
    for name, parser, model in P.SHEETS:
        recs, errs = parser(sheets[name], SOURCE)
        parsed[model] = recs
        errors += errs
        print(f"  {name}: {len(recs)} record(s), {len(errs)} problem(s)")
    return parsed, errors


def main(path: str, dry_run: bool) -> int:
    parsed, errors = parse_all(read_sheets(path))
    if errors:
        print("\nNothing was written. Fix these first:")
        for e in errors:
            print("  -", e)
        return 1
    absent = P.segments_missing_maturity(parsed["CandidateSegment"])
    if absent:
        print("\nNote: the source has no 'Professional Maturity' cell for", "; ".join(absent), "-> stored empty, not guessed.")
    if dry_run:
        print("\nDry run: nothing written.")
        return 0

    from app.database import SessionLocal
    from app.models import reference_data as M

    db = SessionLocal()
    try:
        for model_name, recs in parsed.items():
            model, key = getattr(M, model_name), KEYS[model_name]
            ins = upd = same = 0
            for rec in recs:
                row = db.query(model).filter(model.source == rec["source"],
                                             *[getattr(model, k) == rec[k] for k in key]).first()
                if row is None:
                    db.add(model(**rec))
                    ins += 1
                    continue
                changed = False
                for k, v in rec.items():
                    if getattr(row, k) != v:
                        setattr(row, k, v)
                        changed = True
                upd += changed
                same += not changed
            print(f"  {model_name}: {ins} new, {upd} updated, {same} unchanged")
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()
    return 0


def verify(path: str) -> int:
    """Independent check: re-read the workbook cell by cell and compare to what is stored."""
    from app.database import SessionLocal
    from app.models import reference_data as M
    from app.modules.market_roles.parser import clean

    sheets = read_sheets(path)
    db = SessionLocal()
    diffs = checked = 0
    try:
        for name, (model_name, attrs) in VERIFY.items():
            model = getattr(M, model_name)
            spec = {"Skill Framework": P.SKILL_FRAMEWORK, "Salary Benchmarks": P.SALARY, "Industry Analysis": P.INDUSTRY,
                    "Candidate Segmentation": P.SEGMENTS}[name]
            rows = [(n, v) for n, v in sheets[name] if n > spec["header_row"] and any(clean(c) for c in v)]
            stored = {r.source_row: r for r in db.query(model).filter(model.source == SOURCE).all()}
            if len(stored) != len(rows):
                print(f"{name}: workbook has {len(rows)} rows, DB has {len(stored)}")
                diffs += 1
            for n, v in rows:
                row = stored.get(n)
                if row is None:
                    print(f"{name} row {n}: missing in DB")
                    diffs += 1
                    continue
                cells = [clean(x) for x in list(v)[1:1 + len(attrs)]]
                if name == "Candidate Segmentation" and len(cells) == 13 and cells[12] == "" and cells[11].isdigit():
                    cells = cells[:9] + [""] + cells[9:12]          # the documented absent-cell rows
                for attr, cell in zip(attrs, cells):
                    checked += 1
                    got = getattr(row, attr)
                    got = "" if got is None else str(got)
                    if got != cell and not (got.replace(".0", "") == cell):
                        print(f"{name} row {n} {attr}: workbook {cell!r} != DB {got!r}")
                        diffs += 1
        eco = {(r.source_row): r for r in db.query(M.UpscEcosystemEntry).filter(M.UpscEcosystemEntry.source == SOURCE).all()}
        ws = sheets["UPSC Ecosystem & Skills"]
        sect_rows = 0
        for n, v in ws:
            if n in eco:
                sect_rows += 1
                heads = P.ECOSYSTEM["sections"][eco[n].section]
                cells = [clean(x) for x in list(v)[1:1 + len(heads)]]
                want = [str(eco[n].data.get(h, "")) for h in heads]      # by header: JSONB does not keep key order
                checked += len(cells)
                if cells != want:
                    print(f"Ecosystem row {n}: workbook {cells} != DB {want}")
                    diffs += 1
        print(f"Ecosystem: {len(eco)} stored rows; {sect_rows} matched to workbook rows")
    finally:
        db.close()
    print(f"{checked} values compared, {diffs} difference(s)")
    return 1 if diffs else 0


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args or not Path(args[0]).exists():
        raise SystemExit("Usage: python -m app.scripts.import_reference_sheets <workbook.xlsx> [--dry-run | --verify]")
    raise SystemExit(verify(args[0]) if "--verify" in sys.argv else main(args[0], "--dry-run" in sys.argv))
