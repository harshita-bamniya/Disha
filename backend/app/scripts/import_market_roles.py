"""Import the UPSC market-roles workbook into the `market_roles` reference catalogue.

    python -m app.scripts.import_market_roles "C:/path/UPSC_Master_Database_v2.xlsx" --dry-run
    python -m app.scripts.import_market_roles "C:/path/UPSC_Master_Database_v2.xlsx"

Needs `openpyxl` (a one-off admin tool, so it is not a runtime dependency: pip install openpyxl).

Strict by design: every row is parsed and validated first. If ANY row has a problem, nothing is written
and every problem is listed with its row, company, role and column. Re-running is safe: rows are matched
on (source, company, role) and updated in place, so the catalogue never gets duplicates.
"""
import sys
from pathlib import Path

from app.modules.market_roles.parser import HEADERS, parse_row

SHEET = "Master Database"
SOURCE = "UPSC_Master_Database_v2"
HEADER_ROW = 2          # row 1 is the group banner (COMPANY INFORMATION, JOB INFORMATION, ...)


def read_rows(path: str) -> list[tuple[int, dict]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"Sheet '{SHEET}' not found. Sheets: {wb.sheetnames}")
    rows = list(wb[SHEET].iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else "" for h in rows[HEADER_ROW - 1]]
    missing = [h for h in HEADERS if h not in headers]
    if missing:
        raise SystemExit(f"Workbook is missing expected columns: {missing}")
    out = []
    for i, values in enumerate(rows[HEADER_ROW:], start=HEADER_ROW + 1):
        if not any(v not in (None, "") for v in values):
            continue
        out.append((i, dict(zip(headers, values))))
    return out


def main(path: str, dry_run: bool) -> int:
    parsed, errors, seen = [], [], {}
    rows = read_rows(path)
    for row_no, raw in rows:
        record, errs = parse_row(raw, SOURCE, row_no)
        errors += errs
        if record:
            key = (record["company_name"].lower(), record["job_role"].lower())
            if key in seen:
                errors.append(f"row {row_no}: duplicate of row {seen[key]} ({record['company_name']} / {record['job_role']})")
            seen[key] = row_no
            parsed.append(record)
    print(f"{len(rows)} rows read, {len(parsed)} valid, {len(errors)} problem(s)")
    if errors:
        print("\nNothing was written. Fix these first:")
        for e in errors:
            print("  -", e)
        return 1

    if dry_run:
        print("\nDry run: nothing written. Example of the assembled description for row", parsed[0]["source_row"], "\n")
        print(parsed[0]["full_description"])
        return 0

    from app.database import SessionLocal
    from app.models.market_role import MarketRole

    db = SessionLocal()
    inserted = updated = unchanged = 0
    try:
        for rec in parsed:
            row = db.query(MarketRole).filter(
                MarketRole.source == rec["source"], MarketRole.company_name == rec["company_name"],
                MarketRole.job_role == rec["job_role"]).first()
            if row is None:
                db.add(MarketRole(**rec))
                inserted += 1
                continue
            changed = False
            for key, value in rec.items():
                if getattr(row, key) != value:
                    setattr(row, key, value)
                    changed = True
            updated += changed
            unchanged += not changed
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()
    print(f"Imported: {inserted} new, {updated} updated, {unchanged} unchanged.")
    return 0


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args or not Path(args[0]).exists():
        raise SystemExit("Usage: python -m app.scripts.import_market_roles <workbook.xlsx> [--dry-run]")
    raise SystemExit(main(args[0], "--dry-run" in sys.argv))
