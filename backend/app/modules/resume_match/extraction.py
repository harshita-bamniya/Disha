"""Text extraction from resume files plus ATS-readability signals.

Unlike the older parse path this reads DOCX tables, headers and footers, never
raises on an unreadable file (it reports `text_extractable=False` so the ATS
score can say so), and measures the things an employer's ATS trips over:
tables, images, multi-column layouts, missing standard headings, mixed date
formats.
"""
from __future__ import annotations

import io
import re

from .schemas import AtsSignals

MAX_BYTES = 5 * 1024 * 1024

_HEADINGS = {
    "experience": r"(?:work |professional |relevant |employment )?(?:experience|history)|employment",
    "education": r"education|academics?|qualifications?",
    "skills": r"(?:technical |key |core )?skills|competenc(?:y|ies)|technologies",
    "summary": r"(?:professional |career )?(?:summary|profile|objective)|about me",
    "projects": r"projects?",
    "certifications": r"certifications?|licen[cs]es?",
}
_HEADING_RE = {k: re.compile(rf"^\s*(?:{v})\s*:?\s*$", re.I) for k, v in _HEADINGS.items()}
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_PHONE = re.compile(r"\d[\d\s()+-]{7,}\d")
_MON_YYYY = re.compile(r"\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+(?:19|20)\d{2}\b", re.I)
_MM_YYYY = re.compile(r"\b(?:0?[1-9]|1[0-2])\s*/\s*(?:19|20)\d{2}\b")
_ISO = re.compile(r"\b(?:19|20)\d{2}-(?:0[1-9]|1[0-2])\b")


def signals_from_text(text: str, *, tables: int = 0, images: int = 0, multi_column: bool = False) -> AtsSignals:
    """Derive readability signals from already-extracted text plus layout counts."""
    sections = sorted({
        key for line in text.splitlines() if len(line) < 60
        for key, rx in _HEADING_RE.items() if rx.match(line)
    })
    styles = sum(1 for rx in (_MON_YYYY, _MM_YYYY, _ISO) if rx.search(text))
    return AtsSignals(
        text_extractable=len(text.strip()) >= 50,
        char_count=len(text.strip()),
        sections_found=sections,
        has_contact_info=bool(_EMAIL.search(text) or _PHONE.search(text)),
        dates_consistent=styles <= 1,
        table_count=tables,
        image_count=images,
        multi_column=multi_column,
    )


def _pdf(data: bytes) -> tuple[str, int, int, bool]:
    import fitz  # PyMuPDF

    text_parts: list[str] = []
    tables = images = 0
    multi = False
    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc:
            text_parts.append(page.get_text())
            images += len(page.get_images(full=True))
            try:
                tables += len(page.find_tables().tables)
            except Exception:
                pass
            blocks = [b for b in page.get_text("blocks") if len(b[4].strip()) > 20]
            width = page.rect.width
            left = sum(1 for b in blocks if b[2] < 0.55 * width)
            right = sum(1 for b in blocks if b[0] > 0.45 * width)
            if len(blocks) >= 8 and left >= 0.25 * len(blocks) and right >= 0.25 * len(blocks):
                multi = True
    return "\n".join(text_parts), tables, images, multi


def _docx(data: bytes) -> tuple[str, int, int, bool]:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            seen: list[str] = []
            for cell in row.cells:
                if cell.text and cell.text not in seen:  # merged cells repeat
                    seen.append(cell.text)
            parts.append(" | ".join(seen))
    for section in d.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    multi = any(
        (sect._sectPr.xpath("./w:cols/@w:num") or ["1"])[0] not in ("1", "")
        for sect in d.sections
    )
    return "\n".join(parts), len(d.tables), len(d.inline_shapes), multi


def extract(data: bytes, filename: str) -> tuple[str, AtsSignals]:
    """Return (text, ats_signals). Never raises for an unreadable document."""
    if len(data) > MAX_BYTES:
        raise ValueError("File is larger than 5 MB.")
    try:
        if data[:4] == b"%PDF":
            text, tables, images, multi = _pdf(data)
        elif data[:4] == b"PK\x03\x04" or filename.lower().endswith(".docx"):
            text, tables, images, multi = _docx(data)
        else:
            raise ValueError("Unsupported file type. Upload a PDF or DOCX.")
    except ValueError:
        raise
    except Exception:
        return "", AtsSignals(text_extractable=False)
    return text, signals_from_text(text, tables=tables, images=images, multi_column=multi)
